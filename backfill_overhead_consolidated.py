"""Backfill Overhead and Consolidated EBITDA for 2026 P3-P8 from source workbooks."""
import openpyxl, json, os

BASE = r'C:\Users\ascha\OneDrive\Desktop\forage-data'
JSON_PATH = os.path.join(BASE, 'dashboard_data.json')

# Each entry: (period_str, source_path). Column 2 is always the headline period.
SOURCES = [
    ('3', os.path.join(BASE, 'Kitchen Trailing 12 no OH P3.26.xlsx')),
    ('4', os.path.join(BASE, 'Kitchen Trailing 12 P4.26.xlsx')),
    ('5', os.path.join(BASE, 'Kitchen Trailing 12 P5.26.xlsx')),
    ('6', os.path.join(BASE, 'Kitchen Trailing 12 P6.26 (2).xlsx')),
    ('7', r'C:\Users\ascha\Downloads\Kitchen Trailing P&L P7.26.xlsx'),
    ('8', r'C:\Users\ascha\Downloads\Kitchen Trailing 12 P8.26 (1).xlsx'),
]
COL = 2

def find_row(ws, label):
    for r in range(1, ws.max_row+1):
        v = ws.cell(r,1).value
        if v is None: continue
        if str(v).strip() == label.strip():
            return r
    return None

with open(JSON_PATH) as f:
    data = json.load(f)

for period, src in SOURCES:
    if not os.path.exists(src):
        print(f'SKIP P{period}: source not found: {src}')
        continue
    wb = openpyxl.load_workbook(src, data_only=True)
    for sheet, key in [('Overhead-F', 'Overhead_2026'), ('Consolidated-F', 'Consolidated_2026')]:
        if sheet not in wb.sheetnames:
            print(f'  P{period} {sheet}: sheet missing')
            continue
        ws = wb[sheet]
        row = find_row(ws, 'EBITDA')
        if row is None:
            print(f'  P{period} {sheet}: EBITDA row not found')
            continue
        v = ws.cell(row=row, column=COL).value
        v = round(float(v), 2) if v is not None else 0
        data.setdefault(key, {}).setdefault('EBITDA', {str(p):0 for p in range(1,13)})[period] = v
        print(f'P{period} {key} EBITDA (row {row}) = {v}')
    wb.close()

with open(JSON_PATH, 'w') as f:
    json.dump(data, f, indent=2)
print('Saved', JSON_PATH)
